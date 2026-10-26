"""Independent source-to-API reconciliation using the shipped workbooks."""

from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import math
import unittest

import openpyxl
import test_review_regressions as fixtures


ROOT = Path(__file__).resolve().parents[2]
VALUE_COLUMNS = {"energy": "energy_kwh", "water": "water_kl",
                 "waste": "waste_kg", "solar": "solar_kwh"}


def source_rows(content):
    """Read original cells, independently of the application's pandas importer."""
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    result = {resource: [] for resource in VALUE_COLUMNS}
    for resource, column in VALUE_COLUMNS.items():
        if resource not in workbook.sheetnames:
            continue
        records = workbook[resource].iter_rows(values_only=True)
        headers = [str(value).strip().lower() for value in next(records)]
        for cells in records:
            row = dict(zip(headers, cells))
            if all(value is None for value in cells):
                continue
            value = row.get(column)
            if resource == "water" and value is None and "water_liters" in row:
                value = Decimal(str(row["water_liters"])) / 1000
            date = row["date"]
            if not isinstance(date, datetime):
                date = datetime.fromisoformat(str(date))
            result[resource].append((date, str(row["building"]).strip(), Decimal(str(value))))
    workbook.close()
    return result


def expected_summary(rows, building=None, start=None, end=None):
    selected = {
        resource: [(date, name, value) for date, name, value in records
                   if (not building or name == building)
                   and (not start or date.date() >= start)
                   and (not end or date.date() <= end)]
        for resource, records in rows.items()
    }
    totals = {resource: sum((value for _, _, value in records), Decimal(0))
              for resource, records in selected.items()}
    months = {(date.year, date.month) for records in selected.values() for date, _, _ in records}
    net = max(totals["energy"] - totals["solar"], Decimal(0))
    components = [(Decimal(".5"), net, Decimal(100000), bool(selected["energy"])),
                  (Decimal(".3"), totals["water"], Decimal(20000), bool(selected["water"])),
                  (Decimal(".2"), totals["waste"], Decimal(30000), bool(selected["waste"]))]
    weight_sum = sum((weight for weight, _, _, present in components if present), Decimal(0))
    pressure = sum((weight * min(value / benchmark, Decimal(1))
                    for weight, value, benchmark, present in components if present), Decimal(0))
    return {**totals, "net_energy": net, "carbon": net * Decimal(".82"),
            "gross_carbon": totals["energy"] * Decimal(".82"),
            "solar_avoided_carbon": min(totals["energy"], totals["solar"]) * Decimal(".82"),
            "export_to_grid_kwh": max(totals["solar"] - totals["energy"], Decimal(0)),
            "monthly_energy_cost_rs": net * 8 / max(len(months), 1),
            "green_index": (1 - pressure / weight_sum) * 100 if weight_sum else None,
            "months_covered": max(len(months), 1)}


class DataVerificationTests(unittest.TestCase):
    setUp = fixtures.DashboardRegressionTests.setUp
    tearDown = fixtures.DashboardRegressionTests.tearDown
    admin_headers = fixtures.DashboardRegressionTests.admin_headers

    def get(self, path, params=None):
        response = self.client.get(path, params=params, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def import_content(self, content, name="verification.xlsx"):
        files = {"file": (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        validated = self.client.post("/admin/validate-campus-excel", files=files, headers=self.headers)
        self.assertEqual(validated.status_code, 200, validated.text)
        self.assertTrue(validated.json()["ready"], validated.text)
        response = self.client.post("/admin/upload-campus-excel", files=files, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)

    def assert_summary(self, actual, expected):
        for key, value in expected.items():
            with self.subTest(metric=key):
                if value is None:
                    self.assertIsNone(actual[key])
                else:
                    self.assertAlmostEqual(actual[key], float(value), delta=.011)

    def test_workbooks_scopes_dates_comparisons_and_roundtrip(self):
        self.headers = self.admin_headers()
        for name in ("sample_data/college_campus_realistic_demo.xlsx",
                     "sample_data/campus_resource_dataset_import_ready.xlsx",
                     "sample_data/campus_resource_dataset.xlsx",
                     "green-campus/campus_resource_dataset.xlsx"):
            with self.subTest(workbook=name):
                content = (ROOT / name).read_bytes()
                rows = source_rows(content)
                self.import_content(content, name)
                self.assert_summary(self.get("/dashboard/summary"), expected_summary(rows))
                buildings = sorted({name for records in rows.values() for _, name, _ in records})
                self.assertEqual(self.get("/dashboard/all-buildings"), buildings)
                for building in buildings:
                    self.assert_summary(self.get("/dashboard/summary", {"building": building}),
                                        expected_summary(rows, building))
                dates = sorted({date.date() for records in rows.values() for date, _, _ in records})
                start, end = dates[len(dates) // 2], dates[-1]
                params = {"date_from": start.isoformat(), "date_to": end.isoformat()}
                expected = expected_summary(rows, start=start, end=end)
                self.assert_summary(self.get("/dashboard/summary", params), expected)
                reversed_params = {"date_from": end.isoformat(), "date_to": start.isoformat()}
                self.assert_summary(self.get("/dashboard/summary", reversed_params), expected)
                comparison = self.get("/dashboard/comparison", params)
                prior_end = start - timedelta(days=1)
                prior_start = start - timedelta(days=(end - start).days + 1)
                prior = expected_summary(rows, start=prior_start, end=prior_end)
                self.assert_summary(comparison["previous"], prior)
                for metric, delta in comparison["delta"].items():
                    self.assertAlmostEqual(delta["change"], round(float(expected[metric] or 0), 2)
                                           - round(float(prior[metric] or 0), 2), delta=.011)
                exported = self.client.get("/admin/export-campus-excel", headers=self.headers)
                self.assertEqual(exported.status_code, 200)
                self.assertEqual(source_rows(exported.content), rows)
                self.import_content(exported.content)
                self.assert_summary(self.get("/dashboard/summary"), expected_summary(rows))

    def test_all_modes_history_bounds_rankings_and_bundle_sections(self):
        self.headers = self.admin_headers()
        rows = source_rows((ROOT / "sample_data/college_campus_realistic_demo.xlsx").read_bytes())
        self.import_content((ROOT / "sample_data/college_campus_realistic_demo.xlsx").read_bytes())
        for mode, horizon in (("daily", 14), ("monthly", 12), ("seasonal", 8), ("yearly", 5)):
            with self.subTest(mode=mode):
                bundle = self.get("/dashboard/bundle", {"granularity": mode})
                self.assertEqual(bundle["errors"], {})
                for resource in VALUE_COLUMNS:
                    series = bundle["forecast"][resource]
                    self.assertAlmostEqual(sum(series["historical_values"]),
                                           float(expected_summary(rows)[resource]), delta=.02)
                    self.assertEqual(len(series["forecast_values"]), horizon)
                    for low, estimate, high in zip(series["forecast_lower"], series["forecast_values"], series["forecast_upper"]):
                        self.assertTrue(all(math.isfinite(value) for value in (low, estimate, high)))
                        self.assertTrue(0 <= low <= estimate <= high)
                risk = bundle["risk"][0]
                forecast = bundle["forecast"]
                last_label = forecast["energy"]["forecast_labels"][-1]
                solar_map = dict(zip(forecast["solar"]["forecast_labels"], forecast["solar"]["forecast_values"]))
                expected_net = max(forecast["energy"]["forecast_values"][-1] - solar_map.get(last_label, 0), 0)
                self.assertAlmostEqual(risk["projected_net_energy"], expected_net, delta=.011)
                self.assertAlmostEqual(risk["projected_cost"], expected_net * 8, delta=.011)
                self.assertAlmostEqual(risk["projected_carbon"], expected_net * .82, delta=.011)
                self.assertEqual(len(bundle["performance"]), len(bundle["buildings"]))
        daily = {resource: defaultdict(Decimal) for resource in VALUE_COLUMNS}
        for resource, records in rows.items():
            for date, _, value in records:
                daily[resource][date.date()] += value
        expected_export_days = sum(daily["solar"][day] > value for day, value in daily["energy"].items() if day in daily["solar"])
        self.assertEqual(bundle["alerts"]["export_ready_day_count"], expected_export_days)
        for resource in bundle["data_quality"]["resources"]:
            self.assertEqual(resource["row_count"], len(rows[resource["resource"]]))
            self.assertEqual(resource["covered_months"], 12)
            self.assertEqual(resource["duplicate_readings"], 0)
            self.assertEqual(resource["missing_month_count"], 0)
        ranking = self.get("/dashboard/building-performance")
        maximums = {resource: max(row[resource] for row in ranking) or 1 for resource in ("net_energy", "water", "waste")}
        for index, row in enumerate(ranking, 1):
            self.assertEqual(row["rank"], index)
            self.assert_summary(row, {key: value for key, value in expected_summary(rows, row["building"]).items() if key in row})
            score = sum(weight * row[key] / maximums[key] for key, weight in (("net_energy", .5), ("water", .3), ("waste", .2)))
            self.assertAlmostEqual(row["efficiency_score"], score, delta=.001)
        monthly = {resource: defaultdict(Decimal) for resource in VALUE_COLUMNS}
        for resource, records in rows.items():
            for date, _, value in records:
                monthly[resource][(date.year, date.month)] += value
        for point in self.get("/energy/trend"):
            month, year = map(int, point["month"].split("-"))
            key = year, month
            self.assertAlmostEqual(point["energy"], float(monthly["energy"][key]), delta=.011)
            self.assertAlmostEqual(point["solar"], float(monthly["solar"][key]), delta=.011)
            self.assertAlmostEqual(point["net"], float(max(monthly["energy"][key] - monthly["solar"][key], 0)), delta=.011)
        forecast = self.get("/ai/forecast/resources", {"granularity": "monthly"})
        occupancy = {int(key): value for key, value in forecast["academic_occupancy_factors"].items()}
        for resource in VALUE_COLUMNS:
            series = forecast[resource]
            for label, value in zip(series["forecast_labels"], series["forecast_values"]):
                month = datetime.strptime(label, "%b %Y").month
                history = [float(total) for (_, historical_month), total in monthly[resource].items() if historical_month == month]
                baseline = sum(history) / len(history)
                self.assertAlmostEqual(value, baseline, delta=.011)

    def test_simulation_settings_and_occupancy_dependencies(self):
        self.headers = self.admin_headers()
        content = (ROOT / "sample_data/campus_resource_dataset_import_ready.xlsx").read_bytes()
        self.import_content(content)
        original = self.get("/dashboard/summary")
        for reduction in (0, 25, 100):
            response = self.client.post("/ai/simulate/impact", headers=self.headers, json={
                "energy_reduction_pct": reduction, "water_reduction_pct": reduction,
                "waste_reduction_pct": reduction, "solar_increase_pct": reduction})
            self.assertEqual(response.status_code, 200, response.text)
            data = response.json()
            for key in VALUE_COLUMNS:
                factor = 1 + reduction / 100 if key == "solar" else 1 - reduction / 100
                self.assertAlmostEqual(data["projected"][key], float(expected_summary(source_rows(content))[key]) * factor, delta=.011)
            net = max(data["projected"]["energy"] - data["projected"]["solar"], 0)
            self.assertAlmostEqual(data["projected"]["carbon"], net * .82, delta=.02)
            self.assertAlmostEqual(data["impact"]["cost_saved"], data["baseline"]["energy_cost"] - data["projected"]["energy_cost"], delta=.02)
        before = self.get("/ai/forecast/resources")
        changed = self.client.post("/meta/settings", headers=self.headers, json={
            "metric_overrides": {"energy_cost_per_kwh": 12, "emission_factor_kg_per_kwh": .5},
            "academic_occupancy_factors": {str(month): .2 for month in range(1, 13)}})
        self.assertEqual(changed.status_code, 200, changed.text)
        after = self.get("/dashboard/summary")
        self.assertEqual(after["energy"], original["energy"])
        self.assertAlmostEqual(after["carbon"], after["net_energy"] * .5, delta=.011)
        self.assertAlmostEqual(after["monthly_energy_cost_rs"], original["monthly_energy_cost_rs"] * 1.5, delta=.02)
        self.assertEqual(after["occupancy_average"], .2)
        forecast = self.get("/ai/forecast/resources")
        self.assertEqual(before["solar"]["forecast_values"], forecast["solar"]["forecast_values"])
        self.assertEqual(before["energy"]["forecast_values"], forecast["energy"]["forecast_values"])
        self.assertNotEqual(before["energy"]["occupancy_scenario_values"], forecast["energy"]["occupancy_scenario_values"])
        events = self.get("/meta/audit-trail")["records"]
        self.assertGreaterEqual(len(events), 2)
        action = self.client.post("/meta/settings", headers=self.headers,
                                  json={"action_tracker": {"verify-meter": "completed"}})
        self.assertEqual(action.status_code, 200, action.text)
        self.assertEqual(self.get("/dashboard/summary")["action_tracker"]["verify-meter"], "completed")
        events_after = self.get("/meta/audit-trail")["records"]
        self.assertEqual(len(events_after), len(events) + 1)

    def test_empty_window_does_not_mark_targets_as_achieved(self):
        self.headers = self.admin_headers()
        summary = self.get("/dashboard/summary", {"date_from": "2000-01-01", "date_to": "2000-12-31"})
        self.assertFalse(summary["has_data"])
        for goal in summary["goal_progress"]:
            self.assertEqual(goal["status"], "no_data")
            self.assertEqual(goal["completion_percent"], 0)
        bundle = self.get("/dashboard/bundle")
        self.assertEqual(bundle["data_quality"]["overall_score"], 0)
        self.assertIn("Insufficient Data", [alert["title"] for alert in bundle["alerts"]["alerts"]])
        self.assertNotIn("System Stable", [alert["title"] for alert in bundle["alerts"]["alerts"]])

    def test_rejected_sheet_is_marked_needs_attention(self):
        self.headers = self.admin_headers()
        workbook = openpyxl.load_workbook(ROOT / "sample_data/campus_resource_dataset_import_ready.xlsx")
        workbook["energy"]["C2"] = -5
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        result = self.client.post("/admin/validate-campus-excel", headers=self.headers,
                                  files={"file": ("invalid.xlsx", content.getvalue())})
        self.assertEqual(result.status_code, 200, result.text)
        payload = result.json()
        self.assertFalse(payload["ready"])
        energy = next(sheet for sheet in payload["sheet_summaries"] if sheet["sheet"] == "energy")
        self.assertEqual(energy["status"], "needs_attention")


if __name__ == "__main__":
    unittest.main()
